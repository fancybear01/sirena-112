#!/usr/bin/env python3
"""Проверка форматов обмена и хранения. Задача #98.

Смысл скрипта в том, чтобы матрица форматов не была словами. Каждая строка
в docs/format-matrix.md получена этой проверкой, и её можно повторить.

Что делает:

1. JSON сценариев - примеры из contracts сверяются со схемой.
2. JSON API - отчёт и сводка забираются у Core и разбираются как UTF-8.
3. Круговорот сценария - выгрузили из Core, поменяли id, импортировали,
   прочли обратно и сравнили поле за полем.
4. Отказы импорта - чужая версия классификатора, сломанная схема, дубль id.
   Проверяется не только код ответа, но и что сообщение человеку понятно.
5. XLSX и PDF - файлы скачиваются и открываются настоящими читалками,
   а не проверяются по расширению.
6. WAV - разбирается заголовок и содержимое: настоящий ли это звук.

Запуск (Core должен быть поднят; импорт требует включения и локального адреса):

    python scripts/check_formats.py --core http://127.0.0.1:8080

Код возврата 0 - все проверки прошли, 1 - есть расхождения, они печатаются.
"""

import argparse
import json
import subprocess
import sys
import urllib.error
import urllib.request
import uuid
import wave
from pathlib import Path
from typing import Dict, List, Optional

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = ROOT / "contracts" / "examples"
SCHEMA = ROOT / "contracts" / "scenario.schema.json"
VOICE_FIXTURES = ROOT / "services" / "ai" / "tests" / "fixtures" / "voice"

# Кириллица с «ё» намеренно: на ней ломается и кодировка, и сортировка.
MARKER = "проверка форматов, ёлки-палки"


class Report:
    """Собирает расхождения, чтобы показать их все разом."""

    def __init__(self) -> None:
        self.problems: List[str] = []
        self.notes: List[str] = []

    def check(self, condition: bool, message: str) -> bool:
        if not condition:
            self.problems.append(message)
        return condition

    def note(self, message: str) -> None:
        self.notes.append(message)


def request(url: str, method: str = "GET", body: Optional[dict] = None, raw: bool = False):
    """Запрос к Core. Возвращает код ответа и тело."""
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"} if data else {}
    req = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=30) as response:
            payload = response.read()
            return response.status, payload if raw else json.loads(payload or b"null")
    except urllib.error.HTTPError as error:
        payload = error.read()
        if raw:
            return error.code, payload
        try:
            return error.code, json.loads(payload or b"null")
        except ValueError:
            return error.code, {"raw": payload.decode("utf-8", "replace")}


# Контейнер Core, через который идут запросы импорта, и порт внутри него.
# Импорт разрешён только с локального адреса, а запрос снаружи приходит
# с адреса шлюза Docker и получает 403. Изнутри контейнера адрес настоящий
# loopback, поэтому круговорот проверяется именно так.
_container: Optional[str] = None
_container_port = "8080"


def local_request(path: str, body: dict) -> tuple:
    """POST изнутри контейнера Core, чтобы адрес был локальным."""
    if _container is None:
        return 403, {"message": "не задан --core-container, импорт проверить нельзя"}
    result = subprocess.run(
        ["docker", "exec", "-i", _container, "curl", "-sS", "-X", "POST",
         "http://127.0.0.1:%s%s" % (_container_port, path),
         "-H", "Content-Type: application/json", "--data-binary", "@-",
         "-o", "/dev/stdout", "-w", "\n%{http_code}"],
        input=json.dumps(body), capture_output=True, text=True, timeout=60,
    )
    output = result.stdout.rsplit("\n", 1)
    if len(output) != 2 or not output[1].strip().isdigit():
        return 0, {"message": "не разобрал ответ: %r" % result.stdout[:160]}
    status = int(output[1].strip())
    try:
        return status, json.loads(output[0] or "null")
    except ValueError:
        return status, {"raw": output[0][:200]}


# --- JSON ---------------------------------------------------------------------


def check_scenario_schema(report: Report) -> None:
    print("JSON сценариев против схемы contracts")
    try:
        from jsonschema import Draft202012Validator
    except ImportError:
        report.note("jsonschema не установлен, схема не проверялась")
        print("   пропущено: нет jsonschema")
        return

    validator = Draft202012Validator(json.loads(SCHEMA.read_text(encoding="utf-8")))
    files = sorted(EXAMPLES.glob("scenario-*.json"))
    report.check(bool(files), "в contracts/examples нет сценариев")
    for path in files:
        errors = sorted(validator.iter_errors(json.loads(path.read_text(encoding="utf-8"))), key=str)
        ok = report.check(not errors, "%s не проходит схему: %s" % (
            path.name, errors[0].message if errors else ""))
        print("   %-28s %s" % (path.name, "валиден" if ok else "НЕ ВАЛИДЕН"))


def check_api_json(report: Report, core: str, session_id: Optional[str]) -> None:
    print("\nJSON API: разбирается и в UTF-8")
    targets = [("сводка", "%s/api/teacher/analytics/summary" % core)]
    if session_id:
        targets.append(("отчёт", "%s/api/teacher/sessions/%s/report" % (core, session_id)))

    for title, url in targets:
        status, payload = request(url, raw=True)
        if not report.check(status == 200, "%s: код %s" % (title, status)):
            continue
        try:
            payload.decode("utf-8")
            decoded = True
        except UnicodeDecodeError:
            decoded = False
        report.check(decoded, "%s не в UTF-8" % title)
        body = json.loads(payload)
        print("   %-10s %d байт, ключей %d, только ASCII: %s"
              % (title, len(payload), len(body), payload.isascii()))


# --- круговорот сценария ------------------------------------------------------


def check_round_trip(report: Report, core: str) -> None:
    print("\nКруговорот сценария: выгрузили, импортировали, прочли обратно")
    status, scenarios = request("%s/api/teacher/scenarios" % core)
    if not report.check(status == 200 and scenarios, "Core не отдал сценарии: код %s" % status):
        return

    source = dict(scenarios[0])
    source["id"] = str(uuid.uuid4())
    source["title"] = "%s (%s)" % (source["title"], MARKER)

    status, answer = local_request("/api/teacher/scenarios/import", {"scenarios": [source]})
    if status == 403:
        report.note(
            "Импорт отключён или запрос не с локального адреса: круговорот не проверен. "
            "Нужно CORE_SCENARIO_IMPORT_ENABLED=true и обращение с 127.0.0.1."
        )
        print("   пропущено: импорт отвечает 403 (%s)" % answer.get("message", ""))
        return
    if not report.check(status == 201, "импорт не принял сценарий: %s %s" % (status, answer)):
        return

    status, back = request("%s/api/teacher/scenarios" % core)
    stored = next((item for item in back if item["id"] == source["id"]), None)
    if not report.check(stored is not None, "импортированный сценарий не читается обратно"):
        return

    report.check(MARKER in stored["title"], "кириллица не выжила круговорот: %r" % stored["title"])
    changed = [key for key in source if source[key] != stored.get(key)]
    report.check(not changed, "круговорот изменил поля: %s" % ", ".join(changed))
    print("   импортирован и прочитан, изменённых полей: %s" % (", ".join(changed) or "нет"))
    print("   кириллица с ё: %s" % ("цела" if MARKER in stored["title"] else "ПОБИТА"))


def check_import_refusals(report: Report, core: str) -> None:
    """Отказы обязаны быть понятными человеку, а не кодом без объяснения."""
    print("\nОтказы импорта: понятно ли, что не так")
    status, scenarios = request("%s/api/teacher/scenarios" % core)
    if status != 200 or not scenarios:
        report.note("сценариев нет, отказы импорта не проверялись")
        return

    base = dict(scenarios[0])

    wrong_version = json.loads(json.dumps(base))
    wrong_version["id"] = str(uuid.uuid4())
    wrong_version["groundTruth"]["classifierVersion"] = "999-1999-01-01"

    broken = json.loads(json.dumps(base))
    broken["id"] = str(uuid.uuid4())
    broken.pop("rubric", None)

    cases = [
        ("чужая версия классификатора", wrong_version, "CATALOG_VERSION"),
        ("сломанная схема", broken, "SCHEMA"),
        ("дубль идентификатора", base, "DUPLICATE_ID"),
    ]

    for title, scenario, expected_code in cases:
        status, answer = local_request(
            "/api/teacher/scenarios/import/validate", {"scenarios": [scenario]}
        )
        if status == 403:
            report.note("отказы импорта не проверены: нужен локальный адрес и включённый импорт")
            print("   пропущено: 403")
            return
        codes = [error.get("code") for error in (answer or {}).get("errors", [])]
        messages = [error.get("message", "") for error in (answer or {}).get("errors", [])]
        ok = report.check(expected_code in codes,
                         "%s: ждали код %s, пришло %s" % (title, expected_code, codes))
        human = any(message and not message.isascii() for message in messages)
        report.check(human, "%s: сообщение не на русском: %s" % (title, messages))
        print("   %-28s %s | %s" % (title, expected_code if ok else "НЕ ТОТ КОД",
                                    (messages[0] if messages else "")[:52]))


# --- XLSX и PDF ---------------------------------------------------------------


def check_office_exports(report: Report, core: str, session_id: Optional[str]) -> None:
    print("\nXLSX и PDF: открываются ли настоящими читалками")
    targets = [
        ("summary.xlsx", "%s/api/teacher/analytics/summary.xlsx" % core, b"PK"),
        ("summary.pdf", "%s/api/teacher/analytics/summary.pdf" % core, b"%PDF"),
    ]
    if session_id:
        targets += [
            ("report.xlsx", "%s/api/teacher/sessions/%s/report.xlsx" % (core, session_id), b"PK"),
            ("report.pdf", "%s/api/teacher/sessions/%s/report.pdf" % (core, session_id), b"%PDF"),
        ]

    for name, url, magic in targets:
        status, payload = request(url, raw=True)
        if not report.check(status == 200, "%s: код %s" % (name, status)):
            continue
        report.check(payload.startswith(magic),
                     "%s: не похоже на формат, первые байты %r" % (name, payload[:8]))
        detail = _open_office_file(report, name, payload)
        print("   %-14s %6d байт  %s" % (name, len(payload), detail))


def _open_office_file(report: Report, name: str, payload: bytes) -> str:
    """Открывает файл по-настоящему: XLSX через openpyxl, PDF через pypdf."""
    import io

    if name.endswith(".xlsx"):
        try:
            import warnings

            from openpyxl import load_workbook
        except ImportError:
            report.note("openpyxl не установлен, содержимое XLSX не проверялось")
            return "содержимое не проверялось"
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(io.BytesIO(payload))
        rows = list(workbook.active.iter_rows(values_only=True))
        report.check(len(rows) > 1, "%s: в книге нет строк" % name)
        return "листов %d, строк %d" % (len(workbook.sheetnames), len(rows))

    try:
        from pypdf import PdfReader
    except ImportError:
        report.note("pypdf не установлен, содержимое PDF не проверялось")
        return "содержимое не проверялось"
    reader = PdfReader(io.BytesIO(payload))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    report.check(len(reader.pages) > 0, "%s: в документе нет страниц" % name)
    report.check(bool(text.strip()), "%s: текст не извлекается" % name)
    return "страниц %d, символов %d" % (len(reader.pages), len(text))


# --- звук ---------------------------------------------------------------------


def check_wav(report: Report) -> None:
    print("\nWAV: формат и содержимое")
    files = sorted(VOICE_FIXTURES.glob("*.wav"))
    if not report.check(bool(files), "нет фикстур речи в %s" % VOICE_FIXTURES):
        return

    for path in files:
        with wave.open(str(path), "rb") as reader:
            channels = reader.getnchannels()
            width = reader.getsampwidth()
            rate = reader.getframerate()
            frames = reader.getnframes()
            pcm = reader.readframes(frames)

        report.check((channels, width, rate) == (1, 2, 16000),
                     "%s: не PCM16 mono 16 кГц, а %d кан. %d бит %d Гц"
                     % (path.name, channels, width * 8, rate))
        nonzero = sum(1 for index in range(0, len(pcm), 2) if pcm[index:index + 2] != b"\x00\x00")
        share = nonzero * 100 // max(1, frames)
        report.check(share > 50, "%s: похоже на тишину, ненулевых отсчётов %d%%" % (path.name, share))
        print("   %-28s %d Гц, %.2f с, ненулевых %d%%"
              % (path.name, rate, frames / rate, share))


# --- база ---------------------------------------------------------------------


def check_database(report: Report, container: Optional[str]) -> None:
    print("\nPostgreSQL: кодировка и сортировка")
    if not container:
        report.note("контейнер PostgreSQL не указан, база не проверялась")
        print("   пропущено: не задан --postgres")
        return

    def psql(query: str) -> str:
        result = subprocess.run(
            ["docker", "exec", "-i", container, "psql", "-U", "sirena112", "-d", "sirena112",
             "-tAc", query],
            capture_output=True, text=True, timeout=30,
        )
        return result.stdout.strip()

    encoding = psql("SELECT current_setting('server_encoding')")
    report.check(encoding == "UTF8", "кодировка сервера %r, ожидалась UTF8" % encoding)
    length = psql("SELECT length('задымление, ёлки')")
    report.check(length == "16", "русский текст считается неверно: длина %r вместо 16" % length)
    collate = psql(
        "SELECT datcollate FROM pg_database WHERE datname = current_database()")
    order = psql(
        "WITH w(x) AS (VALUES ('ёлка'),('если'),('дом'),('Яблоко')) "
        "SELECT string_agg(x, '<' ORDER BY x) FROM w")

    print("   кодировка %s, длина русской строки %s, сортировка %s" % (encoding, length, collate))
    print("   порядок слов: %s" % order)
    if not collate.startswith("ru_RU"):
        # Не ошибка сборки, а осознанный риск: отчёты и списки будут
        # сортироваться не по русскому алфавиту.
        report.note(
            "Сортировка базы %s, а не ru_RU: русские слова упорядочиваются не по алфавиту "
            "(%s). Для списков и отчётов это заметно." % (collate, order)
        )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--core", default="http://127.0.0.1:8080")
    parser.add_argument("--session", default=None,
                        help="идентификатор оценённого занятия для отчётов")
    parser.add_argument("--postgres", default=None,
                        help="имя контейнера PostgreSQL для проверки кодировки")
    parser.add_argument("--core-container", default=None,
                        help="имя контейнера Core: импорт разрешён только с локального адреса, "
                             "поэтому эти запросы идут изнутри контейнера")
    parser.add_argument("--core-container-port", default="8080")
    arguments = parser.parse_args()

    global _container, _container_port
    _container = arguments.core_container
    _container_port = arguments.core_container_port

    report = Report()
    print("=" * 78)
    print("Проверка форматов обмена и хранения")
    print("Core: %s" % arguments.core)
    print("=" * 78)

    session_id = arguments.session
    if session_id is None:
        # Занятия берём из истории: списка занятий у Core нет, GET на
        # /api/teacher/sessions отвечает 500 вместо 405, потому что там
        # только POST. Это отдельная находка, см. docs/format-matrix.md.
        status, history = request("%s/api/teacher/history" % arguments.core)
        attempts = (history or {}).get("attempts") or [] if status == 200 else []
        scored = [item for item in attempts if item.get("state") == "SCORED"]
        if scored:
            session_id = scored[0]["sessionId"]
            print("занятие для отчётов: %s" % session_id)
        else:
            report.note("нет оценённого занятия: отчёты XLSX и PDF по занятию не проверялись")

    check_scenario_schema(report)
    check_api_json(report, arguments.core, session_id)
    check_round_trip(report, arguments.core)
    check_import_refusals(report, arguments.core)
    check_office_exports(report, arguments.core, session_id)
    check_wav(report)
    check_database(report, arguments.postgres)

    print("\n" + "=" * 78)
    if report.notes:
        print("Что осталось непроверенным или требует внимания:")
        for note in report.notes:
            print("  -", note)
    if report.problems:
        print("\nРАСХОЖДЕНИЯ:")
        for problem in report.problems:
            print("  -", problem)
        return 1
    print("Все проверенные форматы читаются, пишутся и валидируются.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
