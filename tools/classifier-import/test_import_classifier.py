import json
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import Workbook
from jsonschema import Draft202012Validator, FormatChecker

from import_classifier import ROOT, import_workbook, encoded, normalize, resolve_services
from examples import build_examples, evaluate

CATALOG_PATH = ROOT / "contracts/catalog/classifier-v046-11.json"
CATALOG = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
RECORDS = {r["classifierCode"]: r for r in CATALOG["records"]}


def test_catalog_schema_and_all_references():
    schema = json.loads((ROOT / "contracts/classifier.schema.json").read_text(encoding="utf-8"))
    Draft202012Validator.check_schema(schema)
    Draft202012Validator(schema, format_checker=FormatChecker()).validate(CATALOG)
    service_ids = {s["id"] for s in CATALOG["services"]}
    signs = {s["id"]: s for s in CATALOG["signs"]}
    questions = {q["id"] for q in CATALOG["questions"]}
    columns = {c["sourceColumn"] for c in CATALOG["columns"]}
    assert len(RECORDS) == len(CATALOG["records"]) == 1281
    assert len(signs) == len(CATALOG["signs"])
    assert len(questions) == len(CATALOG["questions"])
    for sign in signs.values():
        assert sign["parentId"] is None or sign["parentId"] in signs
    for column in CATALOG["columns"]:
        assert column["serviceId"] in service_ids
        assert all(t["questionId"] in questions for t in column["condition"]["all"])
    for record in CATALOG["records"]:
        assert set(record["mainServiceIds"]) <= service_ids
        assert all(s is None or s in signs for s in record["signIds"])
        assert set(record["questionIds"]) <= questions
        assert all(r["column"] in columns for r in record["routingRules"])


def test_source_control_totals_and_golden_row():
    records = CATALOG["records"]
    assert sum(r["responseScenarioStatus"] == "MISSING" for r in records) == 21
    assert sum(r["responseScenarioStatus"] == "EXPLICIT_NONE" for r in records) == 9
    assert sum(not r["mainServiceIds"] for r in records) == 102
    negative = [rule for r in records for rule in r["routingRules"] if rule["action"] == "DO_NOT_NOTIFY"]
    assert len(negative) == 741
    assert {r["column"] for r in negative} == {"AA"}
    r = RECORDS["1050602"]
    signs = {s["id"]: s["label"] for s in CATALOG["signs"]}
    assert [signs[s] for s in r["signIds"]] == ["жилой дом", "мусоропровод", "дым"]
    assert r["incidentType"] == "задымление: мусоропровод"
    assert r["ekp35IncidentType"] == "пожар: мусоропровод"
    assert r["responseScenarioCode"] == "1_9"
    assert r["mainServiceIds"] == ["MCHS"]
    assert r["source"]["row"] == 93
    assert any(s["routingNotes"] for s in CATALOG["sections"])


def default_answers():
    return {q["id"]: ("NONE" if q["id"] == "routing.victims-status" else False)
            for q in CATALOG["questions"] if q["inputType"] == "SINGLE_SELECT"}


def test_victims_condition_adds_ambulance_and_negative_branch_is_preserved():
    answers = default_answers()
    r = RECORDS["1050602"]
    initial, unresolved = evaluate(CATALOG, r, answers)
    assert not unresolved
    assert "AMBULANCE" not in {s["id"] for s in initial}
    answers["routing.victims-status"] = "PRESENT"
    routed, unresolved = evaluate(CATALOG, r, answers)
    assert not unresolved
    assert "AMBULANCE" in {s["id"] for s in routed}
    answers["routing.victims-status"] = "NOT_ON_SCENE"
    routed, unresolved = evaluate(CATALOG, r, answers)
    assert not unresolved
    assert "AMBULANCE" not in {s["id"] for s in routed}
    del answers["routing.victims-status"]
    assert evaluate(CATALOG, r, answers)[1]  # unanswered is not False/NONE


def test_ambiguous_source_conditions_cannot_silently_execute():
    answers = default_answers()
    answers["routing.culture-listed-facility"] = True
    assert "classifier.1050602.CE" in evaluate(CATALOG, RECORDS["1050602"], answers)[1]


def test_normalization_aliases_and_multiple_main_services():
    assert normalize("  МОЭСК\u00a0\n  тест ") == "МОЭСК тест"
    assert resolve_services(" Police ,\t AMBULANCE ") == ["AMBULANCE", "POLICE"]
    assert resolve_services("МОСГАЗ") == ["MOSGAZ"]
    assert resolve_services("  ") == []
    with pytest.raises(ValueError, match="Unknown main service"):
        resolve_services("UNKNOWN_SERVICE")
    assert any(r["mainServiceIds"] == ["METRO", "MZD"] for r in CATALOG["records"])


def test_four_scenarios_validate_against_shared_contract_and_match_generated_files():
    schema = json.loads((ROOT / "contracts/scenario.schema.json").read_text(encoding="utf-8"))
    validator = Draft202012Validator(schema, format_checker=FormatChecker())
    examples = build_examples(CATALOG)
    assert {x["category"] for x in examples.values()} == {"FIRE", "ACCIDENT", "MEDICAL", "UTILITY"}
    for code, example in examples.items():
        validator.validate(example)
        assert example == json.loads((ROOT / f"contracts/examples/scenario-{code}.json").read_text(encoding="utf-8"))
    nullable = deepcopy(examples["1050602"])
    truth = nullable["groundTruth"]
    truth.update(responseScenarioCode=None, responseScenarioStatus="MISSING", mainServices=[], requiredServices=[], ekp35IncidentType=None)
    validator.validate(nullable)
    truth["mainService"] = {"id": "MCHS", "displayName": "МЧС"}
    assert list(validator.iter_errors(nullable))  # cannot escape validation through legacy


def make_workbook(tmp_path):
    """Small synthetic XLSX, using the versioned source headers. No private XLSX in CI."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист1"
    # First 14 columns copied from the documented v046 header signature.
    initial = [
        ["Генерация номера", None, None, None, None, None,
         "Блок, как происшествие отображается для специалиста службы 112 (типовые признаки происшествия)",
         None, None, None, None, None, "Сценарий реагирования", "Главная служба"],
        ["Г", "п1", "п2", "п3", "Номер", "Группа происшествий Учет в статистике в службе 112",
         "112 - Признак.1 (тип происшествия)", "112-Признак.2", "112-Признак.3",
         "Дополнительные признаки (не влияют на тип происшествия)", "Итоговый тип происшествия", "ТИП происшествия ЕКП 35", None, None],
        [None] * 14,
    ]
    for row in range(3):
        sheet.append(initial[row] + [c["sourceHeaders"][row] for c in CATALOG["columns"]])
    sheet.append([None] * 4 + [1, "Пожары и задымления"])
    r = RECORDS["1050602"]
    cells = r["codeComponents"] + [1050602] + r["source"]["rawFields"] + [None] * 76
    from openpyxl.utils import column_index_from_string
    for rule in r["routingRules"]:
        cells[column_index_from_string(rule["column"]) - 1] = rule["rawValue"]
    sheet.append(cells)
    path = tmp_path / "classifier.xlsx"
    book.save(path)
    book.close()
    return path


def test_import_repeated_bytes_and_uncached_formula(tmp_path):
    from openpyxl import load_workbook
    path = make_workbook(tmp_path)
    first = import_workbook(path)
    second = import_workbook(path)
    assert encoded(first) == encoded(second)
    assert first[0]["records"][0]["classifierCode"] == "1050602"
    book = load_workbook(path)
    book.active["E5"] = "=A5*1000000+B5*10000+C5*100+D5"
    book.save(path)
    book.close()
    assert import_workbook(path)[0]["records"][0]["classifierCode"] == "1050602"


@pytest.mark.parametrize("cell,value,match", [
    ("M1", "Changed header", "headers changed.*README"),
    ("E5", 9999999, "code differs"),
    ("E5", "=A5+1", "unexpected code formula"),
    ("N5", "ALIEN", "Unknown main service"),
    ("G5", None, "missing incident type or first sign"),
])
def test_invalid_workbook_fails_clearly(tmp_path, cell, value, match):
    from openpyxl import load_workbook
    path = make_workbook(tmp_path)
    book = load_workbook(path)
    book.active[cell] = value
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match=match):
        import_workbook(path)


def test_duplicate_codes_fail(tmp_path):
    from openpyxl import load_workbook
    path = make_workbook(tmp_path)
    book = load_workbook(path)
    sheet = book.active
    sheet.append([c.value for c in sheet[5]])
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="duplicate classifier code"):
        import_workbook(path)
